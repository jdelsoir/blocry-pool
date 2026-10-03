"""Tests for scripts/parse_schedule.py.

Synthetic workbooks are built with openpyxl. Tests against the real workbook
(.local/pool.xlsx) run only when that file is present.
"""

import datetime as dt
import importlib.util
import json
import sys
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "parse_schedule.py"
REAL_XLSX = ROOT / ".local" / "pool.xlsx"

spec = importlib.util.spec_from_file_location("parse_schedule", SCRIPT)
ps = importlib.util.module_from_spec(spec)
sys.modules["parse_schedule"] = ps
spec.loader.exec_module(ps)

DAYS_FR = ["LUNDI", "MARDI", "MERCREDI ", "JEUDI", "VENDREDI", "SAMEDI", "DIMANCHE"]
TIMES = [dt.time(7 + i // 2, 30 * (i % 2)) for i in range(30)]  # 07:00 .. 21:30


def time_row(t: str, header_row=3) -> int:
    h, m = map(int, t.split(":"))
    return header_row + 1 + (h - 7) * 2 + m // 30


def add_week_sheet(wb, name, title=None, title_cell="D1", row2=None, time_col=2,
                   cells=None, merges=None, n_days=7, header_row=3):
    """Create a schedule sheet.

    cells: {(day_index, "HH:MM", "25"|"50"): value}
    merges: [(day_from, day_to, "HH:MM" from, "HH:MM" to, value, cols)] where cols is
            "pair" (25m of first day to 50m of last day) or "25"/"50".
    """
    ws = wb.create_sheet(name)
    first = time_col + 2  # two time columns, then the grid
    if title is not None:
        ws[title_cell] = title
    labels = row2 if row2 is not None else DAYS_FR[:n_days]
    for d in range(n_days):
        c25 = first + 2 * d
        if d < len(labels) and labels[d] is not None:
            ws.cell(header_row - 1, c25, labels[d])
        ws.merge_cells(start_row=header_row - 1, start_column=c25, end_row=header_row - 1, end_column=c25 + 1)
        ws.cell(header_row, c25, "25m")
        ws.cell(header_row, c25 + 1, "50m")
    for i, t in enumerate(TIMES):
        r = header_row + 1 + i
        ws.cell(r, time_col, t)
        ws.cell(r, time_col + 1, (dt.datetime(2000, 1, 1, t.hour, t.minute) + dt.timedelta(minutes=30)).time())
    ws.cell(header_row + 31, first, "*Les changements de configuration sont indicatifs.")
    for (d, t, length), v in (cells or {}).items():
        c = first + 2 * d + (0 if length == "25" else 1)
        ws.cell(time_row(t, header_row), c, v)
    for d1, d2, t1, t2, v, cols in (merges or []):
        r1, r2 = time_row(t1, header_row), time_row(t2, header_row)
        if cols == "pair":
            c1, c2 = first + 2 * d1, first + 2 * d2 + 1
        else:
            off = 0 if cols == "25" else 1
            c1, c2 = first + 2 * d1 + off, first + 2 * d2 + off
        ws.cell(r1, c1, v)
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    return ws


def new_wb():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    return wb


def run(tmp_path, wb, today="2026-10-03", all_weeks=False):
    p = tmp_path / "book.xlsx"
    wb.save(p)
    data, stats = ps.parse_workbook(p, dt.date.fromisoformat(today), all_weeks=all_weeks)
    return data, stats


def week(data, start):
    return next(w for w in data["weeks"] if w["start"] == start)


def slot(w, date, t):
    day = next(d for d in w["days"] if d["date"] == date)
    return next((s for s in day["slots"] if s["t"] == t), None)


def day_of(w, date):
    return next(d for d in w["days"] if d["date"] == date)


# ---------------------------------------------------------------------------
# Cell classification
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("raw,nums,kind,text", [
    (8, [8], None, None),
    (0, [0], None, None),
    (8.0, [8], None, None),
    ("4 Attention fond mobile à 70 cm!", [4], "note", "Attention fond mobile à 70 cm!"),
    ("8(Arrêt de la nage et sortie de l'eau de 8h40 à 9h)", [8], "change",
     "Arrêt de la nage et sortie de l'eau de 8h40 à 9h"),
    ("15         0", [15, 0], None, None),
    ("0          8 (arrêt de la nage et sortie de l'eau de 8h40 à 9h00)", [0, 8], "change",
     "arrêt de la nage et sortie de l'eau de 8h40 à 9h00"),
    ("changement de 8h40 à 9h*", [], "change", "changement de 8h40 à 9h"),
    ("Changement de configuration et sortie de l'eau de 8h45 à 9h", [], "change",
     "Changement de configuration et sortie de l'eau de 8h45 à 9h"),
    ("sortie de l'eau de 18h à 18h20*", [], "change", "sortie de l'eau de 18h à 18h20"),
    ("Changemement de 15h30 à 15h50", [], "change", "Changemement de 15h30 à 15h50"),
    (" changement     18:00-18:10              0         4", [0, 4], "change", "changement 18:00-18:10"),
    ("FERMÉ SAMEDI ET DIMANCHE", [], "closed", "FERMÉ SAMEDI ET DIMANCHE"),
    ("Fermeture exceptionnelle à 17h30", [], "closed", "Fermeture exceptionnelle à 17h30"),
    ("fermeture à 17h45", [], "closed", "fermeture à 17h45"),
    ("1 (à partir de 17h15)", [1], "note", "à partir de 17h15"),
    ("12 changement et sortie de l'eau de 17h30 à 17h40)", [12], "change",
     "changement et sortie de l'eau de 17h30 à 17h40"),
    ("6\n", [6], None, None),
    ("2 !", [2], None, None),
    ("0*", [0], None, None),
    ("\xa0", [], None, None),
    ("changement de configuration pendant 20 min", [], "change",
     "changement de configuration pendant 20 min"),
])
def test_classify_value(raw, nums, kind, text):
    c = ps.classify_value(raw)
    assert c["nums"] == nums
    assert c["kind"] == kind
    assert c["text"] == text


def test_classify_bad_values():
    assert ps.classify_value(-1)["bad"]
    assert ps.classify_value(2.5)["bad"]
    assert ps.classify_value("#REF!")["bad"]
    assert ps.classify_value(None) == {"nums": [], "kind": None, "text": None, "bad": None}


# ---------------------------------------------------------------------------
# Title parsing
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("text,expected", [
    ("Horaire du 5 au 11 octobre", (5, None, 11, 10)),
    ("Horaire du 29 sept au 5 octobre", (29, 9, 5, 10)),
    ("Horaire du 27 au 2 novembre", (27, None, 2, 11)),
    ("Horaire du 23 février au 1 er mars", (23, 2, 1, 3)),
    ("Horaire du 1er au 7 décembre", (1, None, 7, 12)),
    ("Nombre de couloirs de nage (du 21/10 au 27/10) ", (21, 10, 27, 10)),
    ("Semaine du3 au 9 mars", (3, None, 9, 3)),
    ("Semaine du 28 juillet au 3 aout", (28, 7, 3, 8)),
    ("Semaine du 6 octobre au 12 oct ", (6, 10, 12, 10)),
    ("Horaire du 22 au 28 decembre", (22, None, 28, 12)),
    ("Horaire du 9 au 15 fevrier", (9, None, 15, 2)),
    ("Horaire du 7 au 13", (7, None, 13, None)),
])
def test_parse_title(text, expected):
    p = ps.parse_title(text)
    assert (p["sd"], p["sm"], p["ed"], p["em"]) == expected


@pytest.mark.parametrize("text", ["1ere semaine Paques", "Semaines SCOLAIRES (général)", "", None])
def test_parse_title_none(text):
    assert ps.parse_title(text) is None


# ---------------------------------------------------------------------------
# Whole-workbook behaviour on synthetic sheets
# ---------------------------------------------------------------------------

def test_cell_kinds_and_merges(tmp_path):
    wb = new_wb()
    add_week_sheet(
        wb, "Semaine du 28 sept au 4 octobre", title="Horaire du 28 sept au 4 octobre", title_cell="A1",
        cells={
            (0, "07:00", "25"): 0, (0, "07:00", "50"): 8,
            (0, "09:00", "25"): "8 Attention fond mobile à 70 cm!", (0, "09:00", "50"): 0,
            (0, "12:00", "25"): 19, (0, "12:00", "50"): 0,
            (0, "19:00", "25"): None, (0, "19:00", "50"): 3,
            (1, "07:00", "25"): 0, (1, "07:00", "50"): 8,
            (1, "12:00", "25"): 14,  # empty 50m column counts as 0
            (2, "10:00", "25"): 20, (2, "10:00", "50"): 1,  # ceil(20/2)+1 = 11 > 10
            (2, "11:00", "25"): 19, (2, "11:00", "50"): 0,  # 10 lanes, fine
            (2, "17:00", "25"): 4, (2, "17:00", "50"): 1,
        },
        merges=[
            (0, 4, "08:30", "08:30", "changement de 8h40 à 9h*", "pair"),  # banner Mon..Fri
            (3, 3, "12:00", "12:00", "15         0", "pair"),             # two numbers over one pair
            (2, 2, "17:30", "18:00", "Fermeture à 18h", "pair"),          # slot-level closure
        ],
    )
    data, _ = run(tmp_path, wb, today="2026-10-03")
    w = week(data, "2026-09-28")
    assert w["title"] == "Horaire du 28 sept au 4 octobre"
    assert w["sheet"] == "Semaine du 28 sept au 4 octobre"
    assert [d["date"] for d in w["days"]] == [f"2026-{m:02d}-{d:02d}" for m, d in
                                              [(9, 28), (9, 29), (9, 30), (10, 1), (10, 2), (10, 3), (10, 4)]]
    assert slot(w, "2026-09-28", "07:00") == {"t": "07:00", "l25": 0, "l50": 8}
    assert slot(w, "2026-09-28", "09:00") == {"t": "09:00", "l25": 8, "l50": 0, "note": "Attention fond mobile à 70 cm!"}
    assert slot(w, "2026-09-28", "19:00") == {"t": "19:00", "l25": 0, "l50": 3}
    assert slot(w, "2026-09-28", "07:30") is None  # empty half-hour omitted
    for date in ["2026-09-28", "2026-09-29", "2026-09-30", "2026-10-01", "2026-10-02"]:
        assert slot(w, date, "08:30") == {"t": "08:30", "l25": 0, "l50": 0, "change": "changement de 8h40 à 9h"}
    assert slot(w, "2026-10-03", "08:30") is None  # banner stops on Friday
    assert slot(w, "2026-09-29", "12:00") == {"t": "12:00", "l25": 14, "l50": 0}
    assert slot(w, "2026-10-01", "12:00") == {"t": "12:00", "l25": 15, "l50": 0}
    sus = slot(w, "2026-09-30", "10:00")
    assert sus["suspect"] is True
    assert "suspect" not in slot(w, "2026-09-30", "11:00")
    assert any(x["date"] == "2026-09-30" and x["t"] == "10:00" for x in data["warnings"])
    assert slot(w, "2026-09-30", "17:30") == {"t": "17:30", "closed": "Fermeture à 18h"}
    assert slot(w, "2026-09-30", "18:00") == {"t": "18:00", "closed": "Fermeture à 18h"}
    assert day_of(w, "2026-09-30")["closed"] is None
    assert day_of(w, "2026-10-04") == {"date": "2026-10-04", "closed": None, "slots": []}


def test_closure_banner_closes_whole_days(tmp_path):
    wb = new_wb()
    add_week_sheet(
        wb, " Semaine du 5 au 11 octobre", title="Horaire du 5 au 11 octobre", title_cell="A1",
        cells={(0, "07:00", "25"): 0, (0, "07:00", "50"): 8},
        merges=[(5, 6, "07:00", "07:30", "FERMÉ SAMEDI ET DIMANCHE", "pair")],
    )
    data, _ = run(tmp_path, wb, today="2026-10-03")
    w = week(data, "2026-10-05")
    assert w["sheet"] == "Semaine du 5 au 11 octobre"
    for date in ("2026-10-10", "2026-10-11"):
        assert day_of(w, date) == {"date": date, "closed": "FERMÉ SAMEDI ET DIMANCHE", "slots": []}


def test_year_inference_across_year_boundary(tmp_path):
    wb = new_wb()
    for name in ["Semaine du 15 au 21 décembre", "Semaine du 22 au 28 décembre",
                 "Semaine du 29 au 4 janvier", "Semaine du 5 au 11 janvier"]:
        add_week_sheet(wb, name, title=name.replace("Semaine", "Horaire"),
                       cells={(0, "07:00", "50"): 8})
    data, stats = run(tmp_path, wb, today="2026-01-07")
    assert [w["start"] for w in data["weeks"]] == ["2025-12-29", "2026-01-05"]
    data, _ = run(tmp_path, wb, today="2026-01-07", all_weeks=True)
    assert [w["start"] for w in data["weeks"]] == ["2025-12-15", "2025-12-22", "2025-12-29", "2026-01-05"]
    assert week(data, "2025-12-29")["days"][6]["date"] == "2026-01-04"


def test_year_inference_explicit_months_across_boundary(tmp_path):
    wb = new_wb()
    add_week_sheet(wb, "Semaine du 29 au 4 janvier", title="Horaire du 29 décembre au 4 janvier",
                   cells={(0, "07:00", "50"): 8})
    data, _ = run(tmp_path, wb, today="2026-01-02")
    assert [w["start"] for w in data["weeks"]] == ["2025-12-29"]


def test_month_rollover_titles(tmp_path):
    wb = new_wb()
    add_week_sheet(wb, "Semaine du 20 au 26 octobre", title="Horaire du 20 au 26 octobre",
                   cells={(0, "07:00", "50"): 8})
    add_week_sheet(wb, "Semaine du 27 au 2 novembre", title="Horaire du 27 au 2 novembre",
                   cells={(0, "07:00", "50"): 8})
    add_week_sheet(wb, "Semaine du 3 au 9 nov", title="Horaire du 3 au 9 novembre",
                   cells={(0, "07:00", "50"): 8})
    data, _ = run(tmp_path, wb, today="2025-11-05", all_weeks=True)
    assert [w["start"] for w in data["weeks"]] == ["2025-10-20", "2025-10-27", "2025-11-03"]
    assert not any("disagrees" in x["msg"] for x in data["warnings"])


def test_month_rollover_wrong_end_month(tmp_path):
    """'du 30 au 6 juin' really means 30 June to 6 July (the month is the start's)."""
    wb = new_wb()
    add_week_sheet(wb, "Semaine du 23 au 29 juin", title="Horaire du 23 au 29 juin",
                   cells={(0, "07:00", "50"): 8})
    add_week_sheet(wb, "Semaine du 30 au 6 juin", title=None, cells={(0, "07:00", "50"): 8})
    add_week_sheet(wb, "Semaine du 7 au 13 juillet", title="Horaire du 7 au 13 juillet",
                   cells={(0, "07:00", "50"): 8})
    data, _ = run(tmp_path, wb, today="2025-07-08", all_weeks=True)
    assert [w["start"] for w in data["weeks"]] == ["2025-06-23", "2025-06-30", "2025-07-07"]
    w = week(data, "2025-06-30")
    assert w["days"][6]["date"] == "2025-07-06"
    assert any(x["sheet"] == "Semaine du 30 au 6 juin" and "disagrees" in x["msg"] for x in data["warnings"])


def test_title_falls_back_to_sheet_name_month(tmp_path):
    wb = new_wb()
    add_week_sheet(wb, "Semaine du 7 au 13 sept", title="Horaire du 7 au 13", title_cell="A1",
                   cells={(0, "07:00", "50"): 8})
    data, _ = run(tmp_path, wb, today="2026-09-09")
    assert [w["start"] for w in data["weeks"]] == ["2026-09-07"]
    assert data["weeks"][0]["title"] == "Horaire du 7 au 13"


def test_end_date_mismatch_warns_and_trusts_start(tmp_path):
    wb = new_wb()
    add_week_sheet(wb, "Semaine du 14 au 20 sept", title="Horaire du 14 au 21 sept",
                   cells={(0, "07:00", "50"): 8})
    data, _ = run(tmp_path, wb, today="2026-09-15")
    assert data["weeks"][0]["start"] == "2026-09-14"
    assert any("disagrees" in x["msg"] for x in data["warnings"])


def test_row2_datetimes_win_and_old_layout(tmp_path):
    """Oldest layout: times in columns A/B, grid from C, real datetimes in row 2."""
    wb = new_wb()
    row2 = [dt.datetime(2024, 10, 21) + dt.timedelta(days=i) for i in range(7)]
    add_week_sheet(wb, "1ere semaine de TOUSSAINT", title="Nombre de couloirs de nage (du 21/10 au 27/10) ",
                   title_cell="C1", row2=row2, time_col=1,
                   cells={(0, "07:00", "50"): 8, (2, "07:00", "25"): 16})
    # Partial week: Wednesday to Sunday only, no title.
    row2b = [dt.datetime(2024, 10, 30) + dt.timedelta(days=i) for i in range(5)]
    add_week_sheet(wb, "2e semaine de TOUSSAINT", row2=row2b, time_col=1, n_days=5,
                   cells={(0, "07:00", "25"): 16})
    data, stats = run(tmp_path, wb, today="2026-10-03", all_weeks=True)
    assert [w["start"] for w in data["weeks"]] == ["2024-10-21", "2024-10-28"]
    w1 = week(data, "2024-10-21")
    assert slot(w1, "2024-10-21", "07:00") == {"t": "07:00", "l25": 0, "l50": 8}
    assert slot(w1, "2024-10-23", "07:00") == {"t": "07:00", "l25": 16, "l50": 0}
    w2 = week(data, "2024-10-28")
    assert len(w2["days"]) == 7
    assert day_of(w2, "2024-10-28")["slots"] == []
    assert slot(w2, "2024-10-30", "07:00") == {"t": "07:00", "l25": 16, "l50": 0}


def test_row2_day_numbers_infer_month_from_neighbours(tmp_path):
    wb = new_wb()
    add_week_sheet(wb, "Semaine du 9 au 15 décembre", title="Horaire du 9 au 15 décembre",
                   cells={(0, "07:00", "50"): 8})
    add_week_sheet(wb, "1ere semaine vac d hiver", title="Nombre de couloirs de nage \n1ERE SEMAINE VACANCES HIVER",
                   row2=["LUNDI 23", "MARDI 24", "FÉRIÉ ", "JEUDI 26", "VENDREDI 27", "SAMEDI 28", "DIMANCHE 29"],
                   cells={(0, "07:00", "50"): 8})
    add_week_sheet(wb, "Semaine du 6 au 12 janvier", title="Horaire du 6 au 12 janvier",
                   cells={(0, "07:00", "50"): 8})
    data, _ = run(tmp_path, wb, today="2025-01-08", all_weeks=True)
    assert [w["start"] for w in data["weeks"]] == ["2024-12-09", "2024-12-23", "2025-01-06"]


def test_junk_and_undated_sheets_are_skipped(tmp_path):
    wb = new_wb()
    wb.create_sheet("Feuil1")
    add_week_sheet(wb, "Semaines SCOLAIRES (général)", title="Nombre de couloirs de nage (semaines scolaires)",
                   cells={(0, "07:00", "50"): 8})
    ws = wb.create_sheet("Notes")
    ws["A1"] = "du texte sans grille"
    add_week_sheet(wb, "Semaine du 28 sept au 4 octobre", title="Horaire du 28 sept au 4 octobre",
                   cells={(0, "07:00", "50"): 8})
    data, stats = run(tmp_path, wb, today="2026-10-03", all_weeks=True)
    assert [w["start"] for w in data["weeks"]] == ["2026-09-28"]
    assert stats["dated"] == 1
    skipped = {n for _, n, _ in stats["skipped"]}
    assert skipped == {"Feuil1", "Semaines SCOLAIRES (général)", "Notes"}
    assert data["source"]["sheetCount"] == 4


def test_window_keeps_previous_current_and_future_weeks(tmp_path):
    wb = new_wb()
    for name in ["Semaine du 14 au 20 sept", "Semaine du 21 au 27 sept",
                 "Semaine du 28 sept au 4 octobre", "Semaine du 5 au 11 octobre"]:
        add_week_sheet(wb, name, title=name.replace("Semaine", "Horaire"), cells={(0, "07:00", "50"): 8})
    data, _ = run(tmp_path, wb, today="2026-10-03")
    assert [w["start"] for w in data["weeks"]] == ["2026-09-21", "2026-09-28", "2026-10-05"]
    for w in data["weeks"]:
        assert dt.date.fromisoformat(w["start"]).weekday() == 0
        assert len(w["days"]) == 7


def _weeks_wb(names):
    wb = new_wb()
    for name in names:
        add_week_sheet(wb, name, title=name.replace("Semaine", "Horaire"), cells={(0, "07:00", "50"): 8})
    return wb


BASE_NAMES = ["Semaine du 14 au 20 sept", "Semaine du 21 au 27 sept",
              "Semaine du 28 sept au 4 octobre", "Semaine du 5 au 11 octobre"]


def test_out_of_order_sheet_does_not_cascade(tmp_path):
    wb = _weeks_wb(BASE_NAMES + ["Semaine du 26 oct au 1er nov", "Semaine du 12 au 18 octobre"])
    data, stats = run(tmp_path, wb, today="2026-10-03")
    assert [w["start"] for w in data["weeks"]] == ["2026-09-21", "2026-09-28", "2026-10-05", "2026-10-12", "2026-10-26"]
    assert [st for _, _, st, _ in stats["starts"]] == [
        "2026-09-14", "2026-09-21", "2026-09-28", "2026-10-05", "2026-10-26", "2026-10-12"]
    assert any("order broken" in w["msg"] for w in data["warnings"])


def test_swapped_last_sheets_keep_both_weeks(tmp_path):
    wb = _weeks_wb(["Semaine du 14 au 20 sept", "Semaine du 21 au 27 sept",
                    "Semaine du 5 au 11 octobre", "Semaine du 28 sept au 4 octobre"])
    data, _ = run(tmp_path, wb, today="2026-10-03")
    assert [w["start"] for w in data["weeks"]] == ["2026-09-21", "2026-09-28", "2026-10-05"]
    assert any("order broken" in w["msg"] for w in data["warnings"])


def test_corrected_copy_appended_keeps_window(tmp_path):
    wb = _weeks_wb(BASE_NAMES + ["Semaine du 28 sept v2"])
    wb.worksheets[-1]["D1"] = "Horaire du 28 sept au 4 octobre"
    data, _ = run(tmp_path, wb, today="2026-10-03")
    assert [w["start"] for w in data["weeks"]] == ["2026-09-21", "2026-09-28", "2026-10-05"]
    msgs = [w["msg"] for w in data["warnings"]]
    assert any("duplicate week" in m for m in msgs) and any("order broken" in m for m in msgs)


def test_start_before_end_day_stays_in_end_month():
    cands = ps.title_candidates(ps.parse_title("Horaire du 5 au 11 octobre"), range(2020, 2028))
    assert all(d.month == 10 for d, _ in cands)


@pytest.mark.parametrize("raw,nums,kind", [
    ("Fermé jusqu'au 12", [], "closed"),
    ("changement de 8h40 à 9", [], "change"),
    ("Fermeture 3", [], "closed"),
])
def test_trailing_number_in_note_is_not_a_lane_count(raw, nums, kind):
    c = ps.classify_value(raw)
    assert c["nums"] == nums and c["kind"] == kind


def test_header_row_and_columns_found_dynamically(tmp_path):
    wb = new_wb()
    add_week_sheet(wb, "Semaine du 5 au 11 octobre", title="Horaire du 5 au 11 octobre",
                   title_cell="C1", header_row=5, time_col=3,
                   cells={(0, "07:00", "25"): 3, (0, "07:00", "50"): 4})
    data, _ = run(tmp_path, wb, today="2026-10-03")
    assert slot(data["weeks"][0], "2026-10-05", "07:00") == {"t": "07:00", "l25": 3, "l50": 4}


def test_cli_writes_json_and_fails_on_empty_window(tmp_path):
    wb = new_wb()
    add_week_sheet(wb, "Semaine du 5 au 11 octobre", title="Horaire du 5 au 11 octobre",
                   cells={(0, "07:00", "50"): 8})
    xlsx = tmp_path / "b.xlsx"
    wb.save(xlsx)
    lm = tmp_path / "lm.txt"
    lm.write_text("Sat, 03 Oct 2026 07:12:06 GMT\n")
    out = tmp_path / "out" / "schedule.json"
    rc = ps.main(["--in", str(xlsx), "--out", str(out), "--today", "2026-10-03", "--last-modified-file", str(lm)])
    assert rc == 0
    data = json.loads(out.read_text(encoding="utf-8"))
    assert data["version"] == 1
    assert data["source"]["lastModified"] == "2026-10-03T07:12:06Z"
    assert data["source"]["url"] == "https://csblocry.be/piscines/"
    assert data["generatedAt"].endswith("Z")
    out2 = tmp_path / "none.json"
    rc = ps.main(["--in", str(xlsx), "--out", str(out2), "--today", "2027-06-01"])
    assert rc == 1
    assert not out2.exists()
    rc = ps.main(["--in", str(tmp_path / "missing.xlsx"), "--out", str(out2)])
    assert rc == 1


def test_bad_sheet_does_not_crash_run(tmp_path, monkeypatch):
    wb = new_wb()
    add_week_sheet(wb, "Semaine du 28 sept au 4 octobre", title="Horaire du 28 sept au 4 octobre",
                   cells={(0, "07:00", "50"): 8})
    add_week_sheet(wb, "Semaine du 5 au 11 octobre", title="Horaire du 5 au 11 octobre",
                   cells={(0, "07:00", "50"): 8})
    real = ps.parse_sheet_grid

    def flaky(ws, grid):
        if "5 au 11" in ws.title:
            raise RuntimeError("boom")
        return real(ws, grid)

    monkeypatch.setattr(ps, "parse_sheet_grid", flaky)
    data, _ = run(tmp_path, wb, today="2026-10-03")
    assert [w["start"] for w in data["weeks"]] == ["2026-09-28"]
    assert any("boom" in x["msg"] for x in data["warnings"])


def test_no_dashes_in_output_strings(tmp_path):
    wb = new_wb()
    add_week_sheet(wb, "Semaine du 5 au 11 octobre", title="Horaire du 5 au 11 octobre",
                   cells={(0, "07:00", "25"): 30, (0, "07:00", "50"): 2})
    data, _ = run(tmp_path, wb, today="2026-10-03")
    text = json.dumps(data, ensure_ascii=False)
    assert chr(0x2014) not in text and chr(0x2013) not in text


# ---------------------------------------------------------------------------
# Real workbook (only when .local/pool.xlsx exists)
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def real():
    if not REAL_XLSX.exists():
        pytest.skip("real workbook .local/pool.xlsx not present")
    data, stats = ps.parse_workbook(REAL_XLSX, dt.date(2026, 10, 3))
    return data, stats


def test_real_window(real):
    data, _ = real
    starts = [w["start"] for w in data["weeks"]]
    assert "2026-09-28" in starts and "2026-10-05" in starts
    assert starts == sorted(starts)
    assert all(dt.date.fromisoformat(s) >= dt.date(2026, 9, 21) for s in starts)


def test_real_known_values(real):
    data, _ = real
    w1 = week(data, "2026-09-28")
    w2 = week(data, "2026-10-05")
    s = slot(w1, "2026-09-28", "07:00")
    assert (s["l25"], s["l50"]) == (0, 8)
    assert slot(w1, "2026-09-30", "12:00")["l25"] == 19
    assert slot(w1, "2026-10-01", "12:00")["l25"] == 14
    for date in ("2026-10-10", "2026-10-11"):
        d = day_of(w2, date)
        assert d["closed"] == "FERMÉ SAMEDI ET DIMANCHE"
        assert d["slots"] == []
    for w in (w1, w2):
        for d in w["days"][:5]:
            s = slot(w, d["date"], "08:30")
            assert s is not None and "change" in s, (d["date"], s)


def test_real_all_weeks_dating(real):
    data, stats = ps.parse_workbook(REAL_XLSX, dt.date(2026, 10, 3), all_weeks=True)
    assert stats["sheets"] >= 76
    assert stats["dated"] >= 72
    starts = [w["start"] for w in data["weeks"]]
    assert len(starts) == len(set(starts))
    assert all(dt.date.fromisoformat(s).weekday() == 0 for s in starts)
    # Sheet order is chronological.
    ordered = [st for _, _, st, _ in stats["starts"]]
    assert ordered == sorted(ordered)
