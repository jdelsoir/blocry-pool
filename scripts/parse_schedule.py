#!/usr/bin/env python3
"""Parse the Blocry pool lane workbook into site/data/schedule.json.

The workbook has one sheet per week. Each sheet holds a grid: a header row with
25m / 50m column pairs (Monday to Sunday), two time columns on the left (from,
to) and one row per half-hour. No year appears anywhere, so week dates are
inferred from titles, sheet names, row-2 labels and neighbouring sheets.

See docs/SPEC.md ("Data contract" and "Week date inference").
"""

from __future__ import annotations

import argparse
import datetime as dt
import email.utils
import json
import math
import re
import sys
import unicodedata
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.stderr.write("openpyxl is required: pip install openpyxl\n")
    raise

SOURCE_URL = "https://csblocry.be/piscines/"
SCHEMA_VERSION = 1
MAX_LANES = 10
# Year inference (choose_starts): costs are in days.
ORDER_PENALTY = 120      # a sheet that does not start after the previous one
MAX_GAP_DAYS = 120       # larger gaps between consecutive sheets are reported
MAX_AHEAD_DAYS = 92      # a week starting further ahead of today is implausible
MAX_BEHIND_DAYS = 3 * 366  # older than this is increasingly implausible
RECENT_SHEETS = 3        # dating warnings on the last sheets are always reported

DAY_NAMES = {
    "lundi": 0, "mardi": 1, "mercredi": 2, "jeudi": 3,
    "vendredi": 4, "samedi": 5, "dimanche": 6,
}

# Prefix-matched French month names (accents already stripped).
MONTH_PREFIXES = [
    ("jan", 1), ("fev", 2), ("mar", 3), ("avr", 4), ("mai", 5), ("juin", 6),
    ("juil", 7), ("aou", 8), ("sep", 9), ("oct", 10), ("nov", 11), ("dec", 12),
]
MONTH_RE = r"(janv\w*|jan\w*|fev\w*|mars|mar|avr\w*|mai|juin|juil\w*|aou\w*|sep\w*|oct\w*|nov\w*|dec\w*)"

CLOSURE_RE = re.compile(r"\bferm", re.I)
CHANGE_RE = re.compile(r"chang|arr[eê]t|sortie\s+de\s+l", re.I)
# Text ending with a preposition that introduces a time or a date.
TIME_PREP_RE = re.compile(
    r"(?:\b(?:à|a|de|du|des|dès|au|aux|le|vers|jusque|jusqu|entre|et)|['’])\s*$", re.I
)
# A standalone integer: not part of a time (8h40, 18:00, 9H) and not a quantity
# with a unit (70 cm, 20 min).
STANDALONE_INT_RE = re.compile(
    r"(?<![\d:hH.,])(\d{1,3})(?![\d:hH.,])(?!\s*(?:min|cm|m\b|mètres|metres|er\b|h\b))"
)


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def strip_accents(s: str) -> str:
    return "".join(
        ch for ch in unicodedata.normalize("NFD", s) if unicodedata.category(ch) != "Mn"
    )


def collapse(s: str) -> str:
    return re.sub(r"\s+", " ", s.replace("\xa0", " ")).strip()


def monday_of(d: dt.date) -> dt.date:
    return d - dt.timedelta(days=d.weekday())


def month_from_word(word: str | None) -> int | None:
    if not word:
        return None
    w = strip_accents(word.lower())
    for prefix, num in MONTH_PREFIXES:
        if w.startswith(prefix):
            return num
    return None


def as_date(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def as_time(v) -> dt.time | None:
    """Return a time for a time-column cell, or None."""
    if isinstance(v, dt.datetime):
        return v.time()
    if isinstance(v, dt.time):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 0 <= v < 1:
        minutes = round(v * 24 * 60)
        return dt.time(minutes // 60, minutes % 60)
    if isinstance(v, str):
        m = re.fullmatch(r"\s*(\d{1,2})\s*[:hH]\s*(\d{2})?\s*", v)
        if m:
            h, mi = int(m.group(1)), int(m.group(2) or 0)
            if 0 <= h < 24 and 0 <= mi < 60:
                return dt.time(h, mi)
    return None


def is_header_label(v, length: str) -> bool:
    return isinstance(v, str) and re.fullmatch(rf"\s*{length}\s*m\s*", v, re.I) is not None


class Warnings:
    def __init__(self):
        self.items: list[dict] = []

    def add(self, sheet, msg, date=None, t=None, sheet_index=None, surface=None):
        self.items.append({
            "sheet": sheet,
            "date": date.isoformat() if isinstance(date, dt.date) else date,
            "t": t,
            "msg": msg,
            "_idx": sheet_index,
            "_surface": surface,
        })


# ---------------------------------------------------------------------------
# Cell classification
# ---------------------------------------------------------------------------

def clean_text(s: str) -> str:
    """Collapse whitespace, strip footnote stars and wrapping parentheses."""
    s = collapse(s)
    for _ in range(3):
        s = s.strip(" *")
        if s.startswith("(") and (s.endswith(")") or ")" not in s):
            s = s[1:]
        if s.endswith(")") and "(" not in s:
            s = s[:-1]
        s = s.strip(" *")
    return collapse(s)


def has_letters(s: str) -> bool:
    return any(ch.isalpha() for ch in s)


def classify_value(v):
    """Classify one raw cell value.

    Returns a dict: {"nums": [int...], "kind": None|"closed"|"change"|"note",
    "text": str|None, "bad": str|None}.
    """
    out = {"nums": [], "kind": None, "text": None, "bad": None}
    if v is None or isinstance(v, bool):
        return out
    if isinstance(v, (int, float)):
        if isinstance(v, float) and not v.is_integer():
            out["bad"] = f"non-integer lane count {v!r}"
            return out
        n = int(v)
        if n < 0:
            out["bad"] = f"negative lane count {v!r}"
            return out
        out["nums"] = [n]
        return out
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        out["bad"] = f"unexpected date/time value {v!r}"
        return out
    s = str(v).replace("\xa0", " ")
    if not s.strip():
        return out
    if re.fullmatch(r"\s*#[A-Z0-9/]+[!?]?\s*", s):
        out["bad"] = f"spreadsheet error value {s.strip()!r}"
        return out

    # Locate standalone integers, keep only the leading and trailing runs as lane counts.
    matches = list(STANDALONE_INT_RE.finditer(s))
    lead, trail = [], []
    pos = 0
    for m in matches:
        if s[pos:m.start()].strip() == "":
            lead.append(m)
            pos = m.end()
        else:
            break
    end = len(s)
    for m in reversed(matches[len(lead):]):
        # 'changement de 8h40 à 9', 'Fermé jusqu'au 12': a number right after a
        # preposition belongs to the note (a time or a date), not a lane count.
        if re.fullmatch(r"[\s*!.]*", s[m.end():end]) and not TIME_PREP_RE.search(s[:m.start()]):
            trail.insert(0, m)
            end = m.start()
        else:
            break
    if trail and not lead:
        rest = s[:trail[0].start()]
        if CLOSURE_RE.search(rest) and has_letters(rest):
            # A closure note followed by bare numbers: never turn it into an open slot.
            out["bad"] = f"number after a closure note ignored in {collapse(s)!r}"
            trail = []
    nums_m = lead + trail
    text = s
    for m in sorted(nums_m, key=lambda m: m.start(), reverse=True):
        text = text[:m.start()] + " " + text[m.end():]
    out["nums"] = [int(m.group(1)) for m in nums_m]
    text = clean_text(text)
    if text and has_letters(text):
        if CLOSURE_RE.search(text):
            out["kind"] = "closed"
        elif CHANGE_RE.search(text):
            out["kind"] = "change"
        else:
            out["kind"] = "note"
        out["text"] = text
    return out


# ---------------------------------------------------------------------------
# Grid detection
# ---------------------------------------------------------------------------

class Grid:
    def __init__(self, header_row, pairs, time_col, rows):
        self.header_row = header_row  # int
        self.pairs = pairs            # list of (col25, col50)
        self.time_col = time_col      # int, the "from" column
        self.rows = rows              # list of (row, time)


def find_grid(ws) -> Grid | None:
    max_r = min(ws.max_row or 0, 20)
    max_c = min(ws.max_column or 0, 60)
    for r in range(1, max_r + 1):
        labels = []
        for c in range(1, max_c + 1):
            v = ws.cell(r, c).value
            if is_header_label(v, "25"):
                labels.append((c, "25"))
            elif is_header_label(v, "50"):
                labels.append((c, "50"))
        if sum(1 for _, k in labels if k == "25") < 1 or sum(1 for _, k in labels if k == "50") < 1:
            continue
        pairs = []
        i = 0
        while i < len(labels) - 1:
            (c1, k1), (c2, k2) = labels[i], labels[i + 1]
            if k1 != k2 and c2 - c1 <= 2:
                pairs.append((c1, c2) if k1 == "25" else (c2, c1))
                i += 2
            else:
                i += 1
        if not pairs:
            continue
        first_col = min(min(p) for p in pairs)
        # Time column: leftmost column left of the grid holding times below the header.
        time_col = None
        for c in range(1, first_col):
            hits = sum(1 for rr in range(r + 1, r + 6) if as_time(ws.cell(rr, c).value) is not None)
            if hits >= 3:
                time_col = c
                break
        if time_col is None:
            continue
        rows = []
        started = False
        for rr in range(r + 1, r + 60):
            t = as_time(ws.cell(rr, time_col).value)
            if t is None:
                if started:
                    break
                continue
            started = True
            rows.append((rr, t))
        if rows:
            return Grid(r, pairs[:7], time_col, rows)
    return None


def merged_lookup(ws):
    """Map (row, col) -> (min_row, min_col, max_row, max_col) for merged cells."""
    lookup = {}
    for rng in ws.merged_cells.ranges:
        b = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                lookup[(r, c)] = b
    return lookup


def cell_value(ws, merged, r, c):
    b = merged.get((r, c))
    if b:
        return ws.cell(b[0], b[1]).value, b
    return ws.cell(r, c).value, None


# ---------------------------------------------------------------------------
# Week date inference
# ---------------------------------------------------------------------------

TITLE_RE = re.compile(
    r"du\s*(\d{1,2})\s*(?:/\s*(\d{1,2})(?:\s*/\s*(\d{2,4}))?|" + MONTH_RE + r")?\.?\s*"
    r"au\s*(\d{1,2})\s*(?:/\s*(\d{1,2})(?:\s*/\s*(\d{2,4}))?|" + MONTH_RE + r")?"
)


def parse_title(text: str | None):
    """Parse 'du X [mois] au Y [mois]'. Returns dict or None.

    Keys: sd, sm, sy, ed, em, ey (ints or None).
    """
    if not text:
        return None
    s = strip_accents(collapse(str(text)).lower())
    s = re.sub(r"(\d)\s*er\b", r"\1", s)
    m = TITLE_RE.search(s)
    if not m:
        return None
    sd = int(m.group(1))
    sm = int(m.group(2)) if m.group(2) else month_from_word(m.group(4))
    sy = m.group(3)
    ed = int(m.group(5))
    em = int(m.group(6)) if m.group(6) else month_from_word(m.group(8))
    ey = m.group(7)

    def year(y):
        if not y:
            return None
        y = int(y)
        return y + 2000 if y < 100 else y

    res = {"sd": sd, "sm": sm, "sy": year(sy), "ed": ed, "em": em, "ey": year(ey)}
    if not (1 <= sd <= 31 and 1 <= ed <= 31):
        return None
    return res


def info_score(p):
    if not p:
        return -1
    return (p["sm"] is not None) * 2 + (p["em"] is not None) * 2 + (p["sy"] is not None)


def title_candidates(p, years):
    """Candidate (date, priority) list for a parsed title (start must be a Monday)."""
    sd, sm, ed, em = p["sd"], p["sm"], p["ed"], p["em"]
    interps = []  # (month, priority) for the start date
    if sm:
        interps.append((sm, 0))
    elif em:
        prev = 12 if em == 1 else em - 1
        if sd > ed:
            interps += [(prev, 0), (em, 1)]
        else:
            # 'du 5 au 11 octobre': a start day before the end day can only be in the end month.
            interps.append((em, 0))
    else:
        interps += [(mo, 0) for mo in range(1, 13)]
    yrs = [p["sy"]] if p.get("sy") else years
    out = []
    for mo, prio in interps:
        for y in yrs:
            # A start in December with an end month of January belongs to the previous year;
            # all years are enumerated anyway, the Monday filter and neighbours decide.
            try:
                d = dt.date(y, mo, sd)
            except ValueError:
                continue
            if d.weekday() == 0:
                out.append((d, prio))
    return out


def day_label_info(ws, grid, merged):
    """Inspect row above the header: datetimes and day names / numbers per pair."""
    r = grid.header_row - 1
    info = []
    if r < 1:
        return [{"date": None, "weekday": None, "daynum": None} for _ in grid.pairs]
    for c25, c50 in grid.pairs:
        d = None
        wd = None
        dn = None
        for c in (min(c25, c50), max(c25, c50)):
            v, _ = cell_value(ws, merged, r, c)
            if v is None:
                continue
            if as_date(v):
                d = d or as_date(v)
            elif isinstance(v, str):
                s = strip_accents(collapse(v).lower())
                for name, idx in DAY_NAMES.items():
                    if s.startswith(name):
                        wd = idx
                        m = re.search(r"\b(\d{1,2})\b", s)
                        if m:
                            dn = int(m.group(1))
                        break
        info.append({"date": d, "weekday": wd, "daynum": dn})
    return info


def sheet_title_text(ws, grid):
    """Return (title_for_output, list of candidate strings to parse)."""
    texts = []
    last_row = grid.header_row - 2 if grid else 2
    for r in range(1, max(1, last_row) + 1):
        for c in range(1, min(ws.max_column or 1, 30) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip():
                texts.append(collapse(v))
    title = None
    for t in texts:
        if parse_title(t):
            title = t
            break
    if title is None and texts:
        title = max(texts, key=len)
    return title, texts


# ---------------------------------------------------------------------------
# Sheet parsing
# ---------------------------------------------------------------------------

def fmt_t(t: dt.time) -> str:
    return f"{t.hour:02d}:{t.minute:02d}"


def build_slot(ws, merged, grid, day_idx, row, t):
    """Return (slot_dict or None, list of warning messages)."""
    c25, c50 = grid.pairs[day_idx]
    v25, b25 = cell_value(ws, merged, row, c25)
    v50, b50 = cell_value(ws, merged, row, c50)
    msgs = []
    l25 = l50 = None
    texts = {"closed": [], "change": [], "note": []}

    def add_text(cls):
        if cls["kind"] and cls["text"] and cls["text"] not in texts[cls["kind"]]:
            texts[cls["kind"]].append(cls["text"])

    pair_merged = b25 is not None and b25 == b50
    if pair_merged:
        cls = classify_value(v25)
        if cls["bad"]:
            msgs.append(cls["bad"])
        add_text(cls)
        nums = cls["nums"]
        if len(nums) >= 2:
            l25, l50 = nums[0], nums[1]
            if len(nums) > 2:
                msgs.append(f"more than two numbers in merged cell {v25!r}")
        elif len(nums) == 1:
            # One number spread over the 25m+50m pair: give it to the length that
            # was in use just before (or just after), else to 25m.
            side = None
            for rr in (row - 1, row + 1):
                if side:
                    break
                a, ba = cell_value(ws, merged, rr, c25)
                b, bb = cell_value(ws, merged, rr, c50)
                if ba is not None and ba == bb:
                    continue
                na = classify_value(a)["nums"]
                nb = classify_value(b)["nums"]
                if na and na[0] > 0 and not (nb and nb[0] > 0):
                    side = "25"
                elif nb and nb[0] > 0 and not (na and na[0] > 0):
                    side = "50"
            if side == "50":
                l25, l50 = 0, nums[0]
            else:
                l25, l50 = nums[0], 0
            if side is None:
                msgs.append(f"single number in a cell merged over 25m and 50m ({v25!r}), length unclear, assigned to 25m")
    else:
        for which, v in (("25", v25), ("50", v50)):
            cls = classify_value(v)
            if cls["bad"]:
                msgs.append(cls["bad"])
            add_text(cls)
            nums = cls["nums"]
            if not nums:
                continue
            if len(nums) >= 2 and which == "25" and classify_value(v50)["nums"] == []:
                l25, l50 = nums[0], nums[1]
                msgs.append(f"two numbers in one 25m cell {v!r}, read as 25m and 50m")
                continue
            if len(nums) >= 2:
                msgs.append(f"several numbers in one cell {v!r}, kept the first")
            if which == "25":
                l25 = nums[0]
            else:
                l50 = nums[0]

    has_nums = l25 is not None or l50 is not None
    if not has_nums and not any(texts.values()):
        return None, msgs
    slot = {"t": fmt_t(t)}
    if texts["closed"] and not (has_nums and ((l25 or 0) + (l50 or 0)) > 0):
        slot["closed"] = " / ".join(texts["closed"])
        return slot, msgs
    slot["l25"] = l25 or 0
    slot["l50"] = l50 or 0
    if texts["change"]:
        slot["change"] = " / ".join(texts["change"])
    notes = texts["note"] + texts["closed"]  # a closure note next to open lanes stays visible
    if notes:
        slot["note"] = " / ".join(notes)
    if math.ceil(slot["l25"] / 2) + slot["l50"] > MAX_LANES:
        slot["suspect"] = True
        msgs.append(
            f"suspect lane counts l25={slot['l25']} l50={slot['l50']} exceed {MAX_LANES} physical lanes"
        )
    return slot, msgs


def parse_sheet_grid(ws, grid):
    merged = merged_lookup(ws)
    days = []
    for i in range(len(grid.pairs)):
        slots = []
        cell_msgs = []
        for row, t in grid.rows:
            slot, msgs = build_slot(ws, merged, grid, i, row, t)
            for m in msgs:
                cell_msgs.append((fmt_t(t), m))
            if slot:
                slots.append(slot)
        closed = None
        if slots and all("closed" in s for s in slots):
            closed = slots[0]["closed"]
            slots = []
        days.append({"closed": closed, "slots": slots, "msgs": cell_msgs})
    return days, merged


# ---------------------------------------------------------------------------
# Workbook parsing
# ---------------------------------------------------------------------------

def parse_workbook(path, today: dt.date, all_weeks=False, last_modified=None, now=None):
    warn = Warnings()
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = []  # per-sheet records
    years = list(range(today.year - 6, today.year + 2))

    for idx, ws in enumerate(wb.worksheets):
        name = ws.title.strip()
        rec = {"idx": idx, "name": name, "ok": False}
        sheets.append(rec)
        try:
            grid = find_grid(ws)
            if grid is None:
                rec["skip"] = "no 25m/50m grid found"
                warn.add(name, "skipped: no 25m/50m grid found", sheet_index=idx)
                continue
            merged = merged_lookup(ws)
            title, texts = sheet_title_text(ws, grid)
            labels = day_label_info(ws, grid, merged)
            rec.update(grid=grid, title=title, labels=labels, ws=ws)

            # Candidate start dates.
            cands = []  # (date, priority)
            source = None
            row2_dates = [l["date"] for l in labels if l["date"]]
            if row2_dates:
                mondays = {monday_of(d) for d in row2_dates}
                if len(mondays) > 1:
                    warn.add(name, f"row 2 dates span several weeks: {sorted(m.isoformat() for m in mondays)}", sheet_index=idx)
                start = min(mondays)
                cands = [(start, 0)]
                source = "row2-dates"
            title_p = None
            for t in texts:
                title_p = parse_title(t)
                if title_p:
                    break
            name_p = parse_title(name)
            parsed = title_p
            if info_score(name_p) > info_score(title_p) and (title_p is None or (title_p["sm"] is None and title_p["em"] is None)):
                parsed = name_p
            rec["parsed"] = parsed
            if not cands and parsed:
                cands = title_candidates(parsed, years)
                source = "title" if parsed is title_p else "sheet-name"
            if not cands:
                daynum = None
                for l in labels:
                    if l["daynum"] and l["weekday"] is not None:
                        # Day number of that weekday, back to Monday.
                        daynum = (l["daynum"], l["weekday"])
                        break
                if daynum:
                    for y in years:
                        for mo in range(1, 13):
                            try:
                                d = dt.date(y, mo, daynum[0])
                            except ValueError:
                                continue
                            if d.weekday() == daynum[1]:
                                cands.append((d - dt.timedelta(days=daynum[1]), 0))
                    source = "row2-day-numbers"
            if row2_dates and parsed:
                tc = {d for d, _ in title_candidates(parsed, years)}
                if cands[0][0] not in tc:
                    warn.add(name, "title disagrees with row 2 dates, row 2 dates win", sheet_index=idx)
            if not cands:
                rec["skip"] = "no date information (title, sheet name, row 2)"
                warn.add(name, "skipped: no usable date in title, sheet name or row 2", sheet_index=idx)
                continue
            rec["cands"] = cands
            rec["source"] = source
            rec["ok"] = True
        except Exception as exc:  # one bad sheet never kills the run
            rec["skip"] = f"parse error: {exc!r}"
            warn.add(name, f"skipped: parse error {exc!r}", sheet_index=idx)

    # Choose years for all sheets at once (see choose_starts): one out-of-order sheet
    # costs one order penalty instead of dragging every earlier sheet into another year.
    datable = [s for s in sheets if s["ok"]]
    choose_starts(datable, today)
    for prev, cur in zip(datable, datable[1:]):
        a, b = prev["start"], cur["start"]
        if b <= a:
            warn.add(cur["name"], f"order broken: inferred start {b} is not after the previous sheet's start {a} ({prev['name']!r})",
                     sheet_index=cur["idx"], surface=(prev["idx"], cur["idx"]))
        elif (b - a).days > MAX_GAP_DAYS:
            warn.add(cur["name"], f"inferred start {b} is far from its neighbours (previous sheet starts {a})",
                     sheet_index=cur["idx"], surface=(prev["idx"], cur["idx"]))
    if datable and datable[-1]["start"] > monday_of(today) + dt.timedelta(days=MAX_AHEAD_DAYS):
        s = datable[-1]
        warn.add(s["name"], f"inferred start {s['start']} is far ahead of today", sheet_index=s["idx"],
                 surface=(s["idx"],))

    # Validate and build weeks.
    weeks_by_start = {}
    for s in datable:
        start = s["start"]
        name = s["name"]
        if start.weekday() != 0:
            warn.add(name, f"inferred start {start} is not a Monday, skipped", sheet_index=s["idx"])
            s["ok"] = False
            s["skip"] = "start not a Monday"
            continue
        p = s.get("parsed")
        if p and s["source"] != "row2-dates":
            end = start + dt.timedelta(days=6)
            if p["ed"] != end.day or (p["em"] and p["em"] != end.month):
                warn.add(name, f"title end date (day {p['ed']}, month {p['em']}) disagrees with start + 6 days ({end}), start trusted", date=start, sheet_index=s["idx"])
        try:
            week = build_week(s, start, warn)
        except Exception as exc:
            warn.add(name, f"skipped: parse error {exc!r}", date=start, sheet_index=s["idx"])
            s["ok"] = False
            s["skip"] = f"parse error: {exc!r}"
            continue
        if start in weeks_by_start:
            warn.add(name, f"duplicate week {start}, keeping this later sheet over {weeks_by_start[start][1]['sheet']!r}", date=start, sheet_index=s["idx"])
        weeks_by_start[start] = (s["idx"], week)

    window_start = monday_of(today) - dt.timedelta(days=7)
    weeks = []
    kept_idx = set()
    for start in sorted(weeks_by_start):
        if all_weeks or start >= window_start:
            idx, week = weeks_by_start[start]
            weeks.append(week)
            kept_idx.add(idx)

    first_kept = min(kept_idx) if kept_idx else len(sheets)
    recent = {s["idx"] for s in datable[-RECENT_SHEETS:]}
    warnings = []
    for w in warn.items:
        i = w.pop("_idx")
        surface = w.pop("_surface")
        # Dating problems near the window are always shown: they can explain missing weeks.
        surfaced = surface is not None and any(j >= first_kept or j in recent for j in surface)
        if all_weeks or i is None or i in kept_idx or (i >= first_kept and not sheets[i]["ok"]) or surfaced:
            warnings.append(w)

    now = now or dt.datetime.now(dt.timezone.utc)
    if last_modified is None and wb.properties.modified:
        lm = wb.properties.modified
        if lm.tzinfo is None:
            lm = lm.replace(tzinfo=dt.timezone.utc)
        last_modified = lm
    data = {
        "version": SCHEMA_VERSION,
        "generatedAt": now.astimezone(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "source": {
            "lastModified": last_modified.astimezone(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ") if last_modified else None,
            "sheetCount": len(wb.worksheets),
            "url": SOURCE_URL,
        },
        "weeks": weeks,
        "warnings": warnings,
    }
    stats = {
        "sheets": len(sheets),
        "dated": sum(1 for s in sheets if s["ok"]),
        "skipped": [(s["idx"], s["name"], s.get("skip")) for s in sheets if not s["ok"]],
        "starts": [(s["idx"], s["name"], s["start"].isoformat(), s.get("source")) for s in sheets if s["ok"]],
    }
    return data, stats


def choose_starts(datable, today):
    """Pick one candidate start per sheet, minimising a global cost (dynamic programming).

    Unary cost: title interpretation priority, distance outside a plausible range
    around today, and for the last sheet a light pull towards the current week.
    Pairwise cost between consecutive sheets: the gap beyond one week, or
    ORDER_PENALTY plus the overlap when a sheet does not start after the previous one.
    Sets s["start"] on every sheet.
    """
    if not datable:
        return
    this_monday = monday_of(today)
    ahead_limit = this_monday + dt.timedelta(days=MAX_AHEAD_DAYS)
    behind_limit = this_monday - dt.timedelta(days=MAX_BEHIND_DAYS)

    def unary(c, last):
        d, prio = c
        cost = prio * 28
        if d > ahead_limit:
            cost += 1000 + (d - ahead_limit).days * 2
        elif d < behind_limit:
            cost += (behind_limit - d).days * 2
        if last:
            cost += abs((d - this_monday).days) * 0.5
        return cost

    def pair(a, b):
        gap = (b - a).days
        return gap - 7 if gap >= 7 else ORDER_PENALTY + abs(gap)

    layers = []
    n = len(datable)
    for i, s in enumerate(datable):
        cands = sorted(set(s["cands"]), key=lambda c: (c[0], c[1]))
        layers.append(cands)
    best = [unary(c, n == 1) for c in layers[0]]
    back = [[None] * len(layers[0])]
    for i in range(1, n):
        cur, ptr = [], []
        for c in layers[i]:
            u = unary(c, i == n - 1)
            j = min(range(len(layers[i - 1])), key=lambda j: (best[j] + pair(layers[i - 1][j][0], c[0]), j))
            cur.append(best[j] + pair(layers[i - 1][j][0], c[0]) + u)
            ptr.append(j)
        best, back = cur, back + [ptr]
    k = min(range(len(best)), key=lambda k: (best[k], abs((layers[-1][k][0] - this_monday).days)))
    for i in range(n - 1, -1, -1):
        datable[i]["start"] = layers[i][k][0]
        if i:
            k = back[i][k]


def build_week(s, start, warn):
    ws, grid, labels = s["ws"], s["grid"], s["labels"]
    days_raw, _ = parse_sheet_grid(ws, grid)
    # Map each pair to a weekday offset.
    offsets = []
    if len(grid.pairs) == 7 and not any(l["date"] for l in labels):
        offsets = list(range(7))
    else:
        for i, l in enumerate(labels):
            if l["date"]:
                offsets.append((l["date"] - start).days)
            elif l["weekday"] is not None:
                offsets.append(l["weekday"])
            else:
                offsets.append(None)
        # Fill gaps positionally from neighbours.
        for i in range(len(offsets)):
            if offsets[i] is None:
                prev = offsets[i - 1] if i > 0 else None
                offsets[i] = (prev + 1) if prev is not None else i
        if len(grid.pairs) != 7:
            warn.add(s["name"], f"grid has {len(grid.pairs)} day columns instead of 7", date=start, sheet_index=s["idx"])
    days = []
    by_offset = {}
    for off, d in zip(offsets, days_raw):
        if 0 <= off <= 6 and off not in by_offset:
            by_offset[off] = d
        else:
            warn.add(s["name"], f"day column with offset {off} ignored", date=start, sheet_index=s["idx"])
    for off in range(7):
        date = start + dt.timedelta(days=off)
        d = by_offset.get(off)
        if d is None:
            days.append({"date": date.isoformat(), "closed": None, "slots": []})
            continue
        for t, m in d["msgs"]:
            warn.add(s["name"], m, date=date, t=t, sheet_index=s["idx"])
        days.append({"date": date.isoformat(), "closed": d["closed"], "slots": d["slots"]})
    return {
        "start": start.isoformat(),
        "title": s["title"] or s["name"],
        "sheet": s["name"],
        "days": days,
    }


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def read_last_modified(path):
    if not path:
        return None
    p = Path(path)
    if not p.exists():
        return None
    text = p.read_text(encoding="utf-8", errors="replace").strip()
    if not text:
        return None
    if ":" in text and text.lower().startswith("last-modified"):
        text = text.split(":", 1)[1].strip()
    try:
        d = email.utils.parsedate_to_datetime(text)
    except (TypeError, ValueError):
        try:
            d = dt.datetime.fromisoformat(text.replace("Z", "+00:00"))
        except ValueError:
            return None
    if d.tzinfo is None:
        d = d.replace(tzinfo=dt.timezone.utc)
    return d


def main(argv=None):
    ap = argparse.ArgumentParser(description="Parse the Blocry pool workbook into schedule.json")
    ap.add_argument("--in", dest="inp", required=True, help="workbook path (.xlsx)")
    ap.add_argument("--out", required=True, help="output JSON path")
    ap.add_argument("--today", help="override today's date (YYYY-MM-DD), default UTC today")
    ap.add_argument("--last-modified-file", "--last-modified", dest="lm_file",
                    help="sidecar file holding the HTTP Last-Modified header")
    ap.add_argument("--all-weeks", action="store_true", help="debug: ignore the date window")
    ap.add_argument("--stats", action="store_true", help="print dating stats to stderr")
    args = ap.parse_args(argv)

    today = dt.date.fromisoformat(args.today) if args.today else dt.datetime.now(dt.timezone.utc).date()
    try:
        data, stats = parse_workbook(args.inp, today, all_weeks=args.all_weeks,
                                     last_modified=read_last_modified(args.lm_file))
    except Exception as exc:
        sys.stderr.write(f"error: cannot read workbook: {exc!r}\n")
        return 1

    if args.stats:
        sys.stderr.write(f"sheets={stats['sheets']} dated={stats['dated']} skipped={len(stats['skipped'])}\n")
        for i, n, why in stats["skipped"]:
            sys.stderr.write(f"  skip [{i}] {n!r}: {why}\n")
        for i, n, st, src in stats["starts"]:
            sys.stderr.write(f"  [{i}] {st} ({src}) {n!r}\n")

    if not data["weeks"]:
        sys.stderr.write("error: no week could be produced for the date window\n")
        return 1
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"wrote {out} with {len(data['weeks'])} weeks, {len(data['warnings'])} warnings")
    return 0


if __name__ == "__main__":
    sys.exit(main())
